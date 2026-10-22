"""PowerPoint and Word exports of an analysis. Same content as the app: avatar, name, title, summary, then messages,
platforms and places, each marked as confirmed by data, calculated, or a suggestion."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DInches
from docx.shared import Pt as DPt
from docx.shared import RGBColor as DRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from ...engine.personas import avatar, summary

BADGE = {"evidence": "Backed by data", "calculated": "From the numbers", "interpretation": "Our suggestion"}
INK, MUTED, ACCENT, SOFT = RGBColor(0x1B, 0x1F, 0x3B), RGBColor(0x6B, 0x72, 0x88), RGBColor(0x4F, 0x6B, 0xED), RGBColor(0xF3, 0xF5, 0xFB)


def _badge(item: dict[str, Any]) -> str:
    return BADGE.get(item.get("basis", "interpretation"), "Our suggestion")


def _market_lines(intake: dict[str, Any], res: dict[str, Any]) -> list[str]:
    """The market in plain words (the summary bullets cover price and audience; this adds competitors and alternatives)."""
    lines: list[str] = []
    anchors = (res.get("anchor_audience") or {}).get("anchors") or []
    if anchors:
        lines.append("Competitors we looked at: " + ", ".join(a["name"] for a in anchors))
    mk = res.get("market") or {}
    if mk.get("need"):
        lines.append(f"The problem your customers are solving: {mk['need']}")
    for tier, title in (("alternative", "What customers use instead"), ("sector", "The wider sector")):
        for row in (mk.get("tiers") or {}).get(tier, []):
            lines.append(f"{title}: {row['label']}: " + ", ".join(b["name"] for b in row["brands"]))
    return lines


def _summary(intake: dict[str, Any], res: dict[str, Any]) -> dict[str, Any]:
    return res.get("executive_summary") or summary.build(intake, res)


# --------------------------------------------------------------------------------------------------- PowerPoint
def _tb(slide, x, y, w, h, text="", size=14, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = text, Pt(size), bold, color
    return tf


def _bullets(tf, rows: list[tuple[str, str]], size=12):
    """rows: (text, badge). First paragraph already exists (empty)."""
    first = True
    for text, badge in rows:
        p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text, r.font.size, r.font.color.rgb = "• " + text, Pt(size), INK
        if badge:
            b = p.add_run()
            b.text, b.font.size, b.font.color.rgb = f"  [{badge}]", Pt(size - 2), MUTED
        p.space_after = Pt(4)


def _bar(slide, prs, title: str, sub: str = ""):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = INK
    shape.line.fill.background()
    _tb(slide, 0.5, 0.18, 10.5, 0.6, title, 24, True, RGBColor(255, 255, 255))
    if sub:
        _tb(slide, 8.3, 0.28, 4.6, 0.4, sub, 12, False, RGBColor(0xC8, 0xCE, 0xE8))


def _section(slide, x, y, w, title: str, rows: list[tuple[str, str]], empty: str, size=12, h=2.2):
    _tb(slide, x, y, w, 0.35, title, 13, True, ACCENT)
    tf = _tb(slide, x, y + 0.35, w, h, "", size)
    _bullets(tf, rows or [(empty, "")], size)


def build_pptx(intake: dict[str, Any], res: dict[str, Any]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    # 1. title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background()
    _tb(s, 0.8, 2.3, 11.5, 1.2, f"Who buys {intake['name']}?", 40, True, RGBColor(255, 255, 255))
    _tb(s, 0.8, 3.5, 11.5, 0.6, f"Your customers in {intake['location']}: who they are, what they like, how to reach them", 18, False, RGBColor(0xC8, 0xCE, 0xE8))
    _tb(s, 0.8, 6.5, 11.5, 0.4, "Customer groups are illustrations, not real people. Each point says whether it is backed by data or is our suggestion.", 11, False, RGBColor(0x9A, 0xA3, 0xC7))
    # 2. summary
    sm = _summary(intake, res)
    s = prs.slides.add_slide(blank)
    _bar(s, prs, "Summary")
    _tb(s, 0.6, 1.1, 12, 0.6, sm["headline"], 20, True)
    tf = _tb(s, 0.6, 1.8, 12, 4.6, "", 15)
    _bullets(tf, [(f"{b['label']}: {b['text']}", "") for b in sm["bullets"]], 15)
    _tb(s, 0.6, 6.7, 12, 0.5, sm["note"], 11, False, MUTED)
    lines = _market_lines(intake, res)
    if lines:
        s = prs.slides.add_slide(blank)
        _bar(s, prs, "The market around your product")
        tf = _tb(s, 0.6, 1.3, 12, 5, "", 16)
        _bullets(tf, [(l, "") for l in lines], 16)
    # 3. personas
    for p in res.get("personas") or []:
        s = prs.slides.add_slide(blank)
        _bar(s, prs, f"{p.get('display_name', '')} — {p.get('title', '')}"[:70], f"Persona {p.get('rank', '')}")
        s.shapes.add_picture(io.BytesIO(avatar.png(p.get("avatar", "n-35-0"), 400)), Inches(0.6), Inches(1.3), Inches(2.2), Inches(2.2))
        _tb(s, 0.6, 3.6, 3.2, 0.4, f"Ages {p.get('age_label', '')}", 14, True)
        _tb(s, 0.6, 4.7, 3.2, 1.2, "Name and picture are for illustration.", 10, False, MUTED)
        _tb(s, 3.3, 1.2, 9.5, 0.9, p.get("summary") or p.get("headline", ""), 16, True)
        _tb(s, 3.3, 2.0, 9.5, 0.8, p.get("who", ""), 12, False, MUTED)
        _section(s, 3.3, 2.9, 4.6, "What drives them", [(m["text"], _badge(m)) for m in p.get("motivations", [])], "Nothing to show", h=1.7)
        _section(s, 8.2, 2.9, 4.6, "What might stop them", [(m["text"], _badge(m)) for m in p.get("objections", [])], "Nothing to show", h=1.7)
        _section(s, 3.3, 5.0, 9.5, "What they like", [(f"{t['name']} ({t.get('kind', '').replace('_', ' ')})", "Backed by data") for t in p.get("tastes", [])[:4]], "Nothing to show", h=1.4)
        # second slide: how to reach them
        s = prs.slides.add_slide(blank)
        _bar(s, prs, f"How to reach {p.get('display_name', '')}", p.get("title", "")[:40])
        _section(s, 0.6, 1.2, 6.2, "Message ideas to test", [(f"{m.get('angle', '')}: {m['text']}" if m.get("angle") else m["text"], _badge(m)) for m in p.get("messages", [])],
                 "Nothing to show", 13, h=2.8)
        _section(s, 7.1, 1.2, 5.6, "Where to find them online", [(x["platform"], _badge(x)) for x in p.get("platforms", [])], "Nothing to show", 13, h=2.0)
        _section(s, 7.1, 3.4, 5.6, "Places to reach them", [(x["place"], "Backed by data") for x in p.get("where_to_reach", [])],
                 "Nothing to show", 13, h=1.8)
        if p.get("events"):
            _section(s, 0.6, 4.4, 6.2, "Event ideas", [(f"{x['idea']}: {x.get('why', '')}", _badge(x)) for x in p.get("events", [])], "Nothing to show", 13, h=1.6)
        if p.get("caveats"):
            _tb(s, 0.6, 6.5, 12, 0.7, "Worth knowing: " + " ".join(p["caveats"])[:300], 10, False, MUTED)
    # last: how to read
    s = prs.slides.add_slide(blank)
    _bar(s, prs, "About this report")
    tf = _tb(s, 0.6, 1.3, 12, 5, "", 14)
    _bullets(tf, [("Backed by data: taste and market data for your location supports it.", ""),
                  ("From the numbers: worked out from your price and public income figures.", ""),
                  ("Our suggestion: our own reasoning from that data. Treat it as an idea to test.", ""),
                  ("Names and pictures are for illustration. Customer groups are not real people.", ""),
                  ("Taste and market data from Qloo. Income data from the World Bank.", "")], 14)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------------- Word
def _dh(doc, text, level=1):
    return doc.add_heading(text, level=level)


def _dbul(doc, rows: list[tuple[str, str]], empty: str):
    for text, badge in rows or [(empty, "")]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        if badge:
            r = p.add_run(f"  [{badge}]")
            r.font.size, r.font.color.rgb = DPt(8.5), DRGB(0x6B, 0x72, 0x88)


def build_docx(intake: dict[str, Any], res: dict[str, Any]) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = DPt(11)
    _dh(doc, f"Who buys {intake['name']}?", 0)
    doc.add_paragraph(f"Your customers in {intake['location']}: who they are, what they like, how to reach them").runs[0].font.color.rgb = DRGB(0x6B, 0x72, 0x88)
    n = doc.add_paragraph("Customer groups are illustrations, not real people. Each point says whether it is backed by data or is our suggestion.")
    n.runs[0].italic = True
    sm = _summary(intake, res)
    _dh(doc, "Summary", 1)
    doc.add_paragraph(sm["headline"]).runs[0].bold = True
    _dbul(doc, [(f"{b['label']}: {b['text']}", "") for b in sm["bullets"]], "")
    doc.add_paragraph(sm["note"]).runs[0].italic = True
    lines = _market_lines(intake, res)
    if lines:
        _dh(doc, "The market around your product", 1)
        _dbul(doc, [(l, "") for l in lines], "")
    for p in res.get("personas") or []:
        doc.add_page_break()
        t = doc.add_table(rows=1, cols=2)
        t.autofit = False
        t.columns[0].width, t.columns[1].width = DInches(1.4), DInches(5.2)
        t.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(avatar.png(p.get("avatar", "n-35-0"), 300)), width=DInches(1.2))
        c = t.cell(0, 1)
        c.paragraphs[0].add_run(f"{p.get('display_name', '')}").bold = True
        c.paragraphs[0].runs[0].font.size = DPt(22)
        c.add_paragraph(p.get("title", "")).runs[0].font.size = DPt(13)
        c.add_paragraph(f"Ages {p.get('age_label', '')}").runs[0].font.size = DPt(10)
        doc.add_paragraph(p.get("summary") or p.get("headline", "")).runs[0].bold = True
        doc.add_paragraph(p.get("who", ""))
        _dh(doc, "What drives them", 2); _dbul(doc, [(m["text"], _badge(m)) for m in p.get("motivations", [])], "Nothing to show")
        _dh(doc, "What might stop them", 2); _dbul(doc, [(m["text"], _badge(m)) for m in p.get("objections", [])], "Nothing to show")
        _dh(doc, "Message ideas to test", 2)
        _dbul(doc, [(f"{m.get('angle', '')}: {m['text']}" if m.get("angle") else m["text"], _badge(m)) for m in p.get("messages", [])], "Nothing to show")
        _dh(doc, "Where to find them online", 2); _dbul(doc, [(x["platform"], _badge(x)) for x in p.get("platforms", [])], "Nothing to show")
        _dh(doc, "Places to reach them", 2)
        _dbul(doc, [(x["place"], "Backed by data") for x in p.get("where_to_reach", [])], "Nothing to show")
        _dh(doc, "What they like", 2); _dbul(doc, [(f"{x['name']} ({x.get('kind', '').replace('_', ' ')})", "Backed by data") for x in p.get("tastes", [])], "Nothing to show")
        if p.get("events"):
            _dh(doc, "Event ideas", 2); _dbul(doc, [(f"{x['idea']}: {x.get('why', '')}", _badge(x)) for x in p.get("events", [])], "Nothing to show")
        for cv in p.get("caveats", []):
            doc.add_paragraph(f"Worth knowing: {cv}").runs[0].font.size = DPt(9)
    doc.add_page_break()
    _dh(doc, "About this report", 1)
    _dbul(doc, [("Backed by data: taste and market data for your location supports it.", ""),
                ("From the numbers: worked out from your price and public income figures.", ""),
                ("Our suggestion: our own reasoning from that data. Treat it as an idea to test.", ""),
                ("Names and pictures are for illustration. Customer groups are not real people.", ""),
                ("Taste and market data from Qloo. Income data from the World Bank.", "")], "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
