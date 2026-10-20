"""API: auth, ownership, validation, background run in mock mode, progress events, edits, exports, avatars."""

import io
import time

import pytest
from fastapi.testclient import TestClient

from discover.api.main import create_app
from discover.api.settings import ApiSettings

INTAKE = {"name": "Zip 5", "description": "A budget smartphone with a large battery for everyday use.", "category": "smartphone",
          "price": 40000, "currency": "NGN", "location": "Lagos", "competitors": ["Tecno"], "max_personas": 3, "product_type": "smartphone"}


def settings(**kw):
    base = dict(env="development", database_url="sqlite://", secret_key="test-secret", cors_origins=["http://localhost:5173"],
                token_minutes=60, max_runs_per_day=10, allow_signup=True, mock_engine=True, runs_dir="/tmp/discover-test-runs", agent_enabled=False, llm_delay=0.0)
    base.update(kw)
    return ApiSettings(**base)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    # a file database: one shared in-memory connection cannot serve concurrent requests and the background run
    with TestClient(create_app(settings(database_url=f"sqlite:///{tmp_path}/api.db"))) as c:
        yield c


def signup(c, email="a@example.com", password="password123"):
    r = c.post("/api/auth/signup", json={"email": email, "password": password, "name": "Ada"})
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def wait(c, h, aid, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        st = c.get(f"/api/analyses/{aid}/events", headers=h).json()["status"]
        if st in ("complete", "failed"):
            return st
        time.sleep(0.2)
    raise AssertionError("analysis did not finish")


def test_health_and_meta(client):
    assert client.get("/health").json() == {"status": "ok"}
    m = client.get("/api/meta").json()
    assert any(c["id"] == "software_saas" for c in m["categories"]) and any(x["name"] == "Nigeria" for x in m["markets"])


def test_signup_login_me_and_duplicate(client):
    h = signup(client)
    assert client.get("/api/auth/me", headers=h).json()["email"] == "a@example.com"
    assert client.post("/api/auth/signup", json={"email": "A@example.com", "password": "password123"}).status_code == 409
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "wrongwrong"}).status_code == 401
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "password123"}).status_code == 200
    assert client.post("/api/auth/signup", json={"email": "b@example.com", "password": "short"}).status_code == 422


def test_endpoints_require_sign_in(client):
    assert client.get("/api/analyses").status_code == 401
    assert client.post("/api/analyses", json={"intake": INTAKE}).status_code == 401
    assert client.get("/api/analyses", headers={"Authorization": "Bearer nope"}).status_code == 401


def test_invalid_intake_is_a_422_with_readable_problems(client):
    h = signup(client)
    bad = dict(INTAKE, category="spaceship", currency="NAIRA")
    r = client.post("/api/analyses", json={"intake": bad}, headers=h)
    assert r.status_code == 422
    r = client.post("/api/analyses", json={"intake": dict(INTAKE, location="Atlantis")}, headers=h)
    assert r.status_code == 422 and "problems" in r.json()["detail"]


def test_full_run_progress_result_edit_and_exports(client):
    h = signup(client)
    r = client.post("/api/analyses", json={"intake": INTAKE}, headers=h)
    assert r.status_code == 202
    aid = r.json()["id"]
    assert wait(client, h, aid) == "complete"
    ev = client.get(f"/api/analyses/{aid}/events", headers=h).json()
    stages = [e["stage"] for e in ev["events"]]
    assert stages[0] == "start" and "affordability" in stages and "personas" in stages and stages[-1] == "done"
    assert client.get(f"/api/analyses/{aid}/events?since={ev['events'][-1]['seq']}", headers=h).json()["events"] == []
    a = client.get(f"/api/analyses/{aid}", headers=h).json()
    ps = a["result"]["personas"]
    assert len(ps) == 3 and ps[0]["display_name"] and ps[0]["avatar"] and a["result"]["synthetic"] is True
    # edit
    r = client.patch(f"/api/analyses/{aid}/personas/0", json={"title": "My own title", "summary": "My own words"}, headers=h)
    assert r.status_code == 200 and r.json()["result"]["personas"][0]["title"] == "My own title"
    assert "title" in r.json()["result"]["personas"][0]["edited_fields"]
    r = client.patch(f"/api/analyses/{aid}/personas/0", json={"title": ""}, headers=h)  # empty resets to the generated text
    assert r.json()["result"]["personas"][0]["title"] != "My own title"
    assert client.patch(f"/api/analyses/{aid}/personas/9", json={"title": "x"}, headers=h).status_code == 404
    # list
    lst = client.get("/api/analyses", headers=h).json()
    assert lst[0]["id"] == aid and lst[0]["persona_count"] == 3 and lst[0]["status"] == "complete"
    # exports
    from pptx import Presentation
    from docx import Document
    r = client.get(f"/api/analyses/{aid}/export?format=pptx", headers=h)
    assert r.status_code == 200 and "presentationml" in r.headers["content-type"]
    prs = Presentation(io.BytesIO(r.content))
    assert len(prs.slides) == 2 + 2 * 3 + 1
    r = client.get(f"/api/analyses/{aid}/export?format=docx", headers=h)
    d = Document(io.BytesIO(r.content))
    assert any(ps[1]["display_name"] in p.text for p in d.paragraphs) or any(ps[1]["display_name"] in c.text for t in d.tables for row in t.rows for c in row.cells)
    assert client.get(f"/api/analyses/{aid}/export?format=json", headers=h).json()["result"]["personas"]
    assert client.get(f"/api/analyses/{aid}/export?format=md", headers=h).status_code == 200
    assert client.get(f"/api/analyses/{aid}/export?format=exe", headers=h).status_code == 422


def test_users_cannot_see_each_others_analyses(client):
    h1, h2 = signup(client, "one@example.com"), signup(client, "two@example.com")
    aid = client.post("/api/analyses", json={"intake": INTAKE}, headers=h1).json()["id"]
    assert client.get(f"/api/analyses/{aid}", headers=h2).status_code == 404
    assert client.get(f"/api/analyses/{aid}/export?format=json", headers=h2).status_code == 404
    assert client.delete(f"/api/analyses/{aid}", headers=h2).status_code == 404
    wait(client, h1, aid)
    assert client.delete(f"/api/analyses/{aid}", headers=h1).status_code == 204
    assert client.get("/api/analyses", headers=h1).json() == []


def test_one_active_run_per_user_and_daily_cap(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with TestClient(create_app(settings(max_runs_per_day=1, database_url=f"sqlite:///{tmp_path}/cap.db"))) as c:
        h = signup(c)
        aid = c.post("/api/analyses", json={"intake": INTAKE}, headers=h).json()["id"]
        assert c.post("/api/analyses", json={"intake": INTAKE}, headers=h).status_code == 429  # still running
        wait(c, h, aid)
        r = c.post("/api/analyses", json={"intake": INTAKE}, headers=h)
        assert r.status_code == 429 and "Daily limit" in r.json()["detail"]


def test_export_before_completion_is_refused_and_failed_runs_are_reported(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from discover.api.models import Analysis
    app = create_app(settings(database_url=f"sqlite:///{tmp_path}/restart.db"))  # a file: the in-memory db is dropped on shutdown
    with TestClient(app) as c:
        h = signup(c)
        async def seed():
            from sqlalchemy import select
            from discover.api.models import User
            async with app.state.session_factory() as db:
                uid = (await db.execute(select(User))).scalars().first().id
                db.add(Analysis(id="x1", user_id=uid, title="t", intake=INTAKE, status="running"))
                await db.commit()
        c.portal.call(seed)
        assert c.get("/api/analyses/x1/export?format=pptx", headers=h).status_code == 409
    # restart: stale running analyses are failed with an explanation
    with TestClient(app) as c2:
        a = c2.get("/api/analyses/x1", headers=h).json()
        assert a["status"] == "failed" and "restarted" in a["error"]


def test_signup_can_be_closed(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with TestClient(create_app(settings(allow_signup=False, database_url=f"sqlite:///{tmp_path}/closed.db"))) as c:
        assert c.post("/api/auth/signup", json={"email": "a@example.com", "password": "password123"}).status_code == 403


def test_avatar_endpoint_serves_svg(client):
    r = client.get("/api/avatars/f-40-2.svg")
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/svg") and r.text.startswith("<svg")
    assert client.get("/api/avatars/" + "x" * 30 + ".svg").status_code == 404


def test_postgres_urls_are_normalised():
    from discover.api.db import normalize_url
    assert normalize_url("postgres://u:p@h/db").startswith("postgresql+asyncpg://")
    assert normalize_url("postgresql://u:p@h/db").startswith("postgresql+asyncpg://")
    assert normalize_url("sqlite:///x.db") == "sqlite+aiosqlite:///x.db"
    # Neon-style libpq options become asyncpg ones
    n = normalize_url("postgresql://u:p@h/db?sslmode=require&channel_binding=require")
    assert "ssl=require" in n and "sslmode" not in n and "channel_binding" not in n


def test_production_requires_a_secret_key(monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.delenv("SECRET_KEY", raising=False)
    with pytest.raises(RuntimeError):
        ApiSettings.from_env()
